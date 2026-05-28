import os
import io
import smtplib
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from datetime import datetime, timedelta

# ── Airtable config ──────────────────────────────────────────────────────────
AIRTABLE_TOKEN = os.environ["AIRTABLE_API_TOKEN"]
BASE_ID        = "appzSljWNDZOf6pIG"
TABLE_ID       = "tblhbYjOAsNPQND3M"

# ── Email config ─────────────────────────────────────────────────────────────
GMAIL_USER        = os.environ["GMAIL_USER"]
GMAIL_APP_PASSWORD = os.environ["GMAIL_APP_PASSWORD"]
RECIPIENT_EMAIL   = "4363773@gmail.com"

# ── Hebrew month names ────────────────────────────────────────────────────────
HEBREW_MONTHS = {
    1: "ינואר", 2: "פברואר", 3: "מרץ",    4: "אפריל",
    5: "מאי",   6: "יוני",   7: "יולי",   8: "אוגוסט",
    9: "ספטמבר",10: "אוקטובר",11: "נובמבר",12: "דצמבר",
}

# ── Grade fields: (airtable_field_id, display_name) ──────────────────────────
GRADE_FIELDS = [
    # הוראה כוללת שנה א
    ("fldLNRMubrq9B0noi", 'הוראה כוללת א - מבוא לתושבע"פ א'),
    ("fld5Tj0d5oEvOxg8D", "הוראה כוללת א - לשון עברית א"),
    ("fldExQua4pq6kwqMT", 'הוראה כוללת א - מבוא לתושב"ע ב'),
    ("fldgIFa61A98niuk8", "הוראה כוללת א - משנה ומפרשיה"),
    ("fldaEaQfezCu2i2sb", "הוראה כוללת א - גמישות פדגוגית"),
    # הוראה כוללת שנה ב
    ("fldAI4ttA8jTvO4fl", "הוראה כוללת ב - תולדות עם ישראל"),
    ("fld7SjBvyO4PdqoNm", "הוראה כוללת ב - לשון עברית א"),
    ("flddFxtlD9TN4f8pR", "הוראה כוללת ב - גיאוגרפיה כללית ב"),
    ("fldFohF3RoqeaVmx6", 'הוראה כוללת ב - זה"ב'),
    ("fldczzobKSKzjDoez", "הוראה כוללת ב - משנה ומפרשיה"),
    # הוראה כוללת שנה ג
    ("fldP3cQk1UTnX4oOs", "הוראה כוללת ג - אוריינות לשונית"),
    ("fld5vhpNawMMqYauC", 'הוראה כוללת ג - מבוא לתושב"ע ב'),
    ("fldSldxLcYAphQsUH", "הוראה כוללת ג - לשון עברית ב"),
    ("fldwjUTwZe74IqODG", "הוראה כוללת ג - משנה ומפרשיה"),
    ("fldwSY3i77YESSU6T", 'הוראה כוללת ג - מתמטיקה בחז"ל'),
    # חינוך מיוחד שנה א
    ("fldTnEhzbMAJI73uy", 'חינוך מיוחד א - מבוא לתושבע"פ א'),
    ("fldciuioNUqnAbJMZ", "חינוך מיוחד א - לשון עברית א"),
    ("fldGwPtvCCYHLvsP1", 'חינוך מיוחד א - מבוא לתושב"ע ב'),
    ("fldsmYPszUiVCS9Zt", "חינוך מיוחד א - משנה ומפרשיה"),
    # חינוך מיוחד שנה ב
    ("fldsuNpEGBDpoM60S", "חינוך מיוחד ב - לשון עברית א"),
    ("fldPCWPx260cNjzkt", "חינוך מיוחד ב - ליקויים חושיים שמיעה"),
    ("fldvbq2iynNSHijVl", 'חינוך מיוחד ב - זה"ב'),
    ("fldqjUsaAJWnmyJr8", "חינוך מיוחד ב - ריתמיקה ותרפיה בתנועה"),
    ("fldRqKdxgvIVofP4w", "חינוך מיוחד ב - מתמטיקה - ראשית הלמידה"),
    ("fldWWU9ngBwpY4ZQJ", "חינוך מיוחד ב - משנה ומפרשיה"),
    # חינוך מיוחד שנה ג
    ("fld3GJfzhqJ0OScG2", "חינוך מיוחד ג - אוריינות לשונית"),
    ("fldhMQfgX7EOrystP", 'חינוך מיוחד ג - מבוא לתושב"ע ב'),
    ("fldB3mIKiC7GaLeUA", "חינוך מיוחד ג - לשון עברית ב"),
    ("fld9jilOAeCGbroKm", "חינוך מיוחד ג - ליקויים חושיים ראייה"),
    ("fldgX2tkHY1BwjB3P", "חינוך מיוחד ג - מתמטיקה - כפל וחילוק"),
    ("fldFFxIS3HipVjNPi", "חינוך מיוחד ג - משנה ומפרשיה"),
    # גיל הרך שנה א
    ("fld1XcdZHUuKE2XHJ", 'גיל הרך א - מבוא לתושבע"פ א'),
    ("fldovq86KNyjm1MpA", "גיל הרך א - לשון עברית א"),
    ("fldKzHoDCgLSxJ0ls", 'גיל הרך א - מבוא לתושב"ע ב'),
    ("fld6Cnl1uW62vVAZq", "גיל הרך א - משנה ומפרשיה"),
    # גיל הרך שנה ב
    ("fld3WutrAIf3tdrzG", "גיל הרך ב - לשון עברית א"),
    ("fldBi7Sh4J8fXgusC", 'גיל הרך ב - זה"ב'),
    ("fldw3kh1dWgRm3NMD", "גיל הרך ב - ריתמיקה ותרפיה בתנועה"),
    ("fldnbAOuJjQxcWyt3", "גיל הרך ב - מתמטיקה - ראשית הלמידה"),
    ("fldmuHDxLeGpibT4H", "גיל הרך ב - משנה ומפרשיה"),
    # גיל הרך שנה ג
    ("fldmMlIFA6sezrbFG", "גיל הרך ג - אוריינות לשונית"),
    ("fldWylO176IO6oGgq", 'גיל הרך ג - מבוא לתושב"ע ב'),
    ("fldGBjJgaX6d5U9iR", "גיל הרך ג - לשון עברית ב"),
    ("fld0cU1USNJSYaa3L", "גיל הרך ג - מתמטיקה - כפל וחילוק"),
    ("fld4rweZtn7ciPbre", "גיל הרך ג - משנה ומפרשיה"),
    # על יסודי שנה א
    ("fldimhz33p5sUQA3S", "על יסודי א - מבוא לפסיכולוגיה"),
    ("fld7DKxwS2RhHq7Lf", 'על יסודי א - מבוא לתושבע"פ א'),
    ("fldyfqmmPQa669g7w", "על יסודי א - לשון עברית א"),
    ("fldbmklelwHZWrFsM", "על יסודי א - גיאוגרפיה כללית ב"),
    ("fldXCsLkozdnAB2qP", 'על יסודי א - זה"ב'),
    ("fldgAPuHh9KttmyGm", "על יסודי א - משנה ומפרשיה"),
    # על יסודי שנה ב
    ("fldKa0dxKGiiGtRUo", "על יסודי ב - תורת החשיבה והלמידה"),
    ("fld4bypYtacy86csR", "על יסודי ב - אוריינות לשונית"),
    ("fldFCllk1GhciXBuC", 'על יסודי ב - מבוא לתושב"ע ב'),
    ("fldgRumv6A3fAd4g7", "על יסודי ב - לשון עברית ב"),
]


def get_previous_month_range():
    today = datetime.today()
    first_of_current = today.replace(day=1)
    last_of_prev = first_of_current - timedelta(days=1)
    first_of_prev = last_of_prev.replace(day=1)
    return first_of_prev, last_of_prev


def fetch_records(start_date, end_date):
    url = f"https://api.airtable.com/v0/{BASE_ID}/{TABLE_ID}"
    headers = {"Authorization": f"Bearer {AIRTABLE_TOKEN}"}
    filter_formula = (
        "AND("
        f"IS_AFTER({{תאריך עדכון ציון}}, '{start_date.strftime('%Y-%m-%dT00:00:00.000Z')}'),"
        f"IS_BEFORE({{תאריך עדכון ציון}}, '{(end_date + timedelta(days=1)).strftime('%Y-%m-%dT00:00:00.000Z')}')"
        ")"
    )
    records, offset = [], None
    while True:
        params = {"filterByFormula": filter_formula}
        if offset:
            params["offset"] = offset
        resp = requests.get(url, headers=headers, params=params)
        resp.raise_for_status()
        data = resp.json()
        records.extend(data.get("records", []))
        offset = data.get("offset")
        if not offset:
            break
    return records


def lookup_val(value):
    if isinstance(value, list):
        return value[0] if value else ""
    return value or ""


def format_date(iso_str):
    if not iso_str:
        return ""
    try:
        return datetime.fromisoformat(iso_str.replace("Z", "+00:00")).strftime("%d/%m/%Y %H:%M")
    except Exception:
        return iso_str


def create_excel(records, month_name, year):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"ציונים {month_name} {year}"
    ws.sheet_view.rightToLeft = True

    hdr_font  = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    hdr_fill  = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alt_fill  = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    data_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    base_headers = ["שם תלמיד", "ת.ז.", "שלוחה", "כיתה", "מסלול", "שנת לימודים", "תאריך עדכון ציון"]

    # Keep only grade columns that have at least one value
    active_grades = [
        (fid, name) for fid, name in GRADE_FIELDS
        if any(r.get("fields", {}).get(fid) for r in records)
    ]

    all_headers = base_headers + [name for _, name in active_grades]

    for col, header in enumerate(all_headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font  = hdr_font
        cell.fill  = hdr_fill
        cell.alignment = hdr_align
    ws.row_dimensions[1].height = 40

    for row, record in enumerate(records, 2):
        f = record.get("fields", {})
        row_data = [
            lookup_val(f.get("fldnPwi0DvosSH8yn")),   # שם ומשפחה
            lookup_val(f.get("fld1WxsDrxku4pvco")),   # ת.ז.
            lookup_val(f.get("fldtqn6x1G8wR5udW")),   # שלוחה
            lookup_val(f.get("fld0o1ZeFzayonZG0")),   # כיתה
            lookup_val(f.get("fldxRxwoTJCowveT4")),   # מסלול
            lookup_val(f.get("fldCjNLzxzAJ7uwWD")),   # שנת לימודים
            format_date(f.get("fldeHCO4yf37oVTwS")),  # תאריך עדכון ציון
        ] + [f.get(fid, "") for fid, _ in active_grades]

        fill = alt_fill if row % 2 == 0 else None
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = data_align
            if fill:
                cell.fill = fill

    for col, header in enumerate(all_headers, 1):
        ws.column_dimensions[get_column_letter(col)].width = min(len(header) + 4, 32)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def send_email(excel_bytes, month_name, year, count):
    msg = MIMEMultipart()
    msg["From"]    = GMAIL_USER
    msg["To"]      = RECIPIENT_EMAIL
    msg["Subject"] = f"דוח ציונים חודשי — {month_name} {year}"

    body = (
        f"שלום,\n\n"
        f"מצורף דוח ציוני תלמידים שעודכנו בחודש {month_name} {year}.\n"
        f'סה"כ רשומות שעודכנו: {count}\n\n'
        f"הדוח הופק אוטומטית."
    )
    msg.attach(MIMEText(body, "plain", "utf-8"))

    part = MIMEBase("application", "octet-stream")
    part.set_payload(excel_bytes)
    encoders.encode_base64(part)
    part.add_header(
        "Content-Disposition",
        f"attachment; filename=grades_{month_name}_{year}.xlsx",
    )
    msg.attach(part)

    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
        server.login(GMAIL_USER, GMAIL_APP_PASSWORD)
        server.sendmail(GMAIL_USER, RECIPIENT_EMAIL, msg.as_string())


def main():
    start, end = get_previous_month_range()
    month_name = HEBREW_MONTHS[start.month]
    year = start.year

    print(f"Fetching grades updated in {month_name} {year} ({start.date()} → {end.date()})...")
    records = fetch_records(start, end)
    print(f"Found {len(records)} updated records.")

    if not records:
        print("No updates this month — no report sent.")
        return

    print("Building Excel...")
    excel_bytes = create_excel(records, month_name, year)

    print("Sending email...")
    send_email(excel_bytes, month_name, year, len(records))
    print("Done! Report sent successfully.")


if __name__ == "__main__":
    main()
